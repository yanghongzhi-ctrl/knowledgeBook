from pathlib import Path
from datetime import date
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "教材型精准RAG知识库项目开发原理_技术方案与建库方法指南.docx"

TITLE = "教材型精准RAG知识库项目开发原理、技术方案与建库方法指南"
SUBTITLE = "以《道路工程数字化设计》AI知识库学习平台为例"
DOC_DATE = "2026年6月"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, east_asia="微软雅黑", latin="Arial", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中右键此处并选择“更新域”以生成目录。"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    heading_settings = {
        "Title": ("微软雅黑", 24, True, (31, 78, 121)),
        "Subtitle": ("微软雅黑", 14, False, (70, 70, 70)),
        "Heading 1": ("微软雅黑", 16, True, (31, 78, 121)),
        "Heading 2": ("微软雅黑", 13, True, (39, 105, 85)),
        "Heading 3": ("微软雅黑", 11, True, (75, 75, 75)),
    }
    for style_name, (font_name, size, bold, color) in heading_settings.items():
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("教材型精准RAG知识库项目开发方法指南")
    set_run_font(run, size=9, color=(110, 110, 110))

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_para(doc, text="", bold_prefix=None, style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF3F0")
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    r = p.add_run(title + "：")
    set_run_font(r, bold=True, color=(39, 105, 85))
    r = p.add_run(text)
    set_run_font(r)
    doc.add_paragraph()


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F7")
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, east_asia="等线", latin="Consolas", size=8.5, color=(45, 55, 60))
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = ""
        set_cell_shading(hdr[i], "1F4E79")
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(header))
        set_run_font(r, bold=True, color=(255, 255, 255), size=9.5)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], "F7FAFC")
            p = cells[i].paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, size=9)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_document():
    doc = Document()
    configure_document(doc)

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    set_run_font(r, east_asia="微软雅黑", size=24, bold=True, color=(31, 78, 121))
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(SUBTITLE)
    set_run_font(r, east_asia="微软雅黑", size=15, color=(70, 70, 70))

    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目方法论与技术实施参考文件")
    set_run_font(r, east_asia="微软雅黑", size=12, bold=True, color=(39, 105, 85))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(DOC_DATE)
    set_run_font(r, east_asia="微软雅黑", size=11, color=(90, 90, 90))

    add_page_break(doc)

    add_heading(doc, "摘要", 1)
    add_para(
        doc,
        "本指南系统总结《道路工程数字化设计》AI知识库学习平台的开发原理、技术方案、研发思路、"
        "知识库建库方法与工程实施细节。项目的核心不是把教材简单切块后交给大模型临时回答，"
        "而是围绕教学目标构建“答案卡优先、证据片段支撑、资源对象扩展、自动评测驱动”的精准RAG系统。"
        "该方法适用于教材、课程、技术规范、企业培训手册和专业知识服务等需要答案稳定、来源可追溯、"
        "内容可维护、效果可量化的项目。"
    )
    add_para(
        doc,
        "文档将本项目的已验证实践与可复用建议分开表达：前者说明当前系统如何落地，后者说明其他项目"
        "如何依据自身内容特征进行裁剪和迁移。"
    )
    add_note(
        doc,
        "核心结论",
        "高质量教材问答系统的竞争力主要来自内容工程、知识对象设计、检索约束和评测闭环，而不是单纯依赖更大的模型。",
    )

    add_heading(doc, "适用范围", 2)
    for item in [
        "面向教材、课程讲义、专业规范、企业知识手册的精准问答系统。",
        "要求回答与正式文稿一致，且需要区分定义、功能、原则、流程、对比等不同教学意图的场景。",
        "包含图片、公式、表格、互动脚本等多模态学习资源的知识库。",
        "需要从单章样板逐步扩展到多章节、多课程或多项目的知识平台。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "文档使用方式", 2)
    add_para(
        doc,
        "项目负责人可重点阅读第1、2、3、13、16章；内容建设人员可重点阅读第4、5、6章；"
        "研发人员可重点阅读第7至12章；验收与运维人员可重点阅读第11、14、15章及附录。"
    )

    add_page_break(doc)
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    add_toc(p)
    add_page_break(doc)

    # 1
    add_heading(doc, "1 项目背景与问题定义", 1)
    add_heading(doc, "1.1 项目目标", 2)
    add_para(
        doc,
        "本项目面向《道路工程数字化设计》教材，建设一个能够帮助学习者准确理解知识点、快速定位教材内容、"
        "调用图表公式与互动资源、并支持持续评测和迭代的AI学习平台。系统需要同时满足教学准确性、工程可维护性"
        "和交互友好性。"
    )
    add_heading(doc, "1.2 普通文档RAG的典型局限", 2)
    for item in [
        "教材切块边界通常服从版式而非教学意图，容易把定义、功能、原则、示例混在一起。",
        "相似问题之间容易互相污染，例如“平面CAD系统功能”与“纵断面CAD系统功能”只差少量词语。",
        "连续追问容易继承错误上下文，例如询问“人机分工的内容”却返回“人机分工的原则”。",
        "模型临时归纳会造成答案结构不稳定、遗漏必答点、混入不应出现的内容。",
        "图片、公式、表格和互动脚本若只作为附件保存，无法被检索、推荐和评测。",
        "没有标准测试集时，系统看似能回答，但无法判断版本升级是否引入回归。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "1.3 本项目的问题重构", 2)
    add_para(
        doc,
        "项目将“让大模型阅读教材并回答问题”重构为“把教材转化为可检索、可验证、可发布的教学知识对象，"
        "再由RAG系统稳定地选择和呈现这些对象”。因此，系统的基本单元不是文档片段，而是围绕学习问题组织的答案卡。"
    )

    # 2
    add_heading(doc, "2 基本原理", 1)
    add_heading(doc, "2.1 答案卡优先，而不是原文切块优先", 2)
    add_para(
        doc,
        "答案卡（Answer Card）是面向一个明确学习问题设计的标准答案对象，包含规范问题、问题类型、答案要点、"
        "必答点、禁答点、输出模式、关联知识点和检索文本。检索系统首先寻找最匹配的答案卡，原文片段只承担证据支撑。"
    )
    add_table(
        doc,
        ["比较维度", "文档切块优先", "答案卡优先"],
        [
            ["检索对象", "段落或固定长度文本", "围绕教学问题组织的标准答案"],
            ["答案稳定性", "依赖模型临时归纳", "由答案要点和输出模式约束"],
            ["相似问题区分", "容易受相邻内容干扰", "通过问题模式、范围词和意图区分"],
            ["质量评测", "难以定义唯一正确目标", "可直接评测目标答案卡与必答点"],
            ["内容维护", "修改原文后难追踪影响", "知识对象可单独版本化与核验"],
        ],
        widths=[3.2, 6.0, 6.0],
    )
    add_heading(doc, "2.2 确定性回答骨架与生成式表达相结合", 2)
    add_para(
        doc,
        "系统不把最终回答完全交给大模型自由生成。答案要点构成确定性骨架，输出模式决定呈现方式，"
        "原文证据和大模型只用于补充说明、语言润色或复杂场景下的解释。这样既保留自然交互能力，又守住教学内容边界。"
    )
    add_heading(doc, "2.3 证据层、答案层与资源层分离", 2)
    add_table(
        doc,
        ["层次", "主要对象", "职责", "禁止事项"],
        [
            ["答案层", "Answer_Cards", "直接回答学习问题", "不得依赖随机生成决定核心结论"],
            ["知识层", "Knowledge_Points、Concept_Comparison", "补充定义、关系和对比", "不得替代明确答案目标"],
            ["证据层", "Source_Chunks", "追溯正式文稿、提供引用依据", "不得直接作为默认答案正文"],
            ["资源层", "Resources、Interactive_Scripts", "提供图、表、公式、互动学习入口", "不得仅作为不可检索附件"],
        ],
        widths=[2.2, 4.2, 5.2, 4.8],
    )
    add_heading(doc, "2.4 评测驱动，而不是主观体验驱动", 2)
    add_para(
        doc,
        "每个版本都应通过自动评测验证：问题是否命中正确答案卡、必答点是否覆盖、禁答内容是否出现、"
        "格式是否符合要求、资源是否推荐正确。用户体验测试仍然重要，但必须建立在可重复的回归测试之上。"
    )
    add_heading(doc, "2.5 保守自动化与人工核验并行", 2)
    add_para(
        doc,
        "Word资源抽取、编号识别、语义绑定和公式结构化可以自动完成大部分工作，但对编号冲突、图题错位、"
        "复杂公式和低置信度匹配应保留人工确认环节。自动化的目标是缩小人工核验范围，而不是掩盖不确定性。"
    )

    # 3
    add_heading(doc, "3 总体技术架构", 1)
    add_heading(doc, "3.1 架构概览", 2)
    add_code_block(
        doc,
        [
            "教材正式稿 Word + 初稿 JSON/JSONL + 图片/公式/表格/互动脚本",
            "                         ↓",
            "内容核验、规范化、资源抽取、知识对象增强",
            "                         ↓",
            "JSON 发布包（可审阅、可版本化、可重建）",
            "                         ↓",
            "PostgreSQL 主数据 + pgvector 检索索引",
            "                         ↓",
            "答案卡检索 / 混合检索 / 资源检索 / RAG 编排 API",
            "                         ↓",
            "学习前端：问题题库、答案、证据、资源、评测与反馈",
            "                         ↘",
            "自动评测与回归报告 → 内容和检索策略持续迭代",
        ],
    )
    add_heading(doc, "3.2 技术栈与职责", 2)
    add_table(
        doc,
        ["技术或组件", "本项目用途", "选择理由"],
        [
            ["Python", "数据处理、校验、检索、评测、API脚本", "标准库能力强，适合内容工程与批处理"],
            ["JSON / JSONL", "维护包与发布包", "易审阅、易比较、易迁移、可重建数据库"],
            ["PostgreSQL", "运行时主数据与评测记录", "结构化查询、事务、索引与版本管理成熟"],
            ["pgvector", "答案卡向量检索", "与结构化数据共库，便于混合检索和运维"],
            ["Ollama + bge-m3", "本地嵌入向量生成", "数据可控，适合离线或内网环境"],
            ["轻量Web API", "问答、资源、详情、评测查询", "接口清晰，便于前后端解耦"],
            ["HTML/CSS/JavaScript", "学习助手前端", "部署简单，便于快速验证交互"],
        ],
        widths=[3.2, 5.5, 6.8],
    )
    add_heading(doc, "3.3 数据源与事实来源原则", 2)
    add_bullet(doc, "Word正式稿是教材内容核验的权威来源。")
    add_bullet(doc, "JSON发布包是内容维护、审阅、版本留档和数据库重建的载体。")
    add_bullet(doc, "PostgreSQL是运行时主数据和评测记录的事实来源。")
    add_bullet(doc, "pgvector索引属于可重建的检索派生数据，不应成为唯一内容来源。")
    add_bullet(doc, "图片、视频和互动HTML保存路径与元数据，数据库通常不存储大体积二进制文件。")

    # 4
    add_heading(doc, "4 数据分层与核心知识对象", 1)
    add_heading(doc, "4.1 核心数据对象", 2)
    add_table(
        doc,
        ["对象", "作用", "关键字段示例"],
        [
            ["Chapter_Structure", "描述章节结构和范围", "chapter_id、section_id、title、order"],
            ["Knowledge_Points", "表示可复用知识单元", "kp_id、title、definition、key_points、status"],
            ["Answer_Cards", "表示可直接回答的问题与标准答案", "answer_card_id、canonical_question、answer_points、must_include、avoid"],
            ["Concept_Comparison", "表示易混概念对比", "comparison_id、concepts、dimensions、conclusion"],
            ["Synonyms_Questions", "扩展学生问法和同义表达", "alias、target_id、priority"],
            ["Knowledge_Relations", "表达知识点关联", "source_id、relation_type、target_id"],
            ["Source_Chunks", "提供正式文稿证据", "chunk_id、source_text、source_location、no_direct_output"],
            ["Resources", "表示图、表、公式和互动资源", "resource_id、type、title、reference_aliases、path、detail"],
            ["Interactive_Scripts", "表示可操作学习步骤", "script_id、title、steps、trigger_questions"],
            ["QA_Evaluation_Testset", "答案卡问答回归集", "question、expected_answer_card_id、expected_points、should_not_include"],
            ["Resource_Evaluation_Testset", "资源导向问答回归集", "question、expected_resource_id、resource_type"],
            ["RAG_Config", "保存检索与输出策略", "threshold、top_k、priority、policy"],
            ["Embedding_Corpus", "保存待嵌入文本定义", "embedding_id、source_type、source_id、text_for_embedding"],
        ],
        widths=[3.6, 5.1, 6.8],
    )
    add_heading(doc, "4.2 答案卡的最低质量标准", 2)
    for item in [
        "一个答案卡只服务一个清晰的问题意图，避免同时回答多个范围不同的问题。",
        "规范问题应使用学生自然语言表达，同时保留教材术语。",
        "答案要点应可独立核验，粒度适合列表、表格、步骤或段落输出。",
        "must_include用于定义不可遗漏的核心内容，avoid用于防止相似概念串答。",
        "answer_mode明确输出形态，例如bullet、table、step、paragraph。",
        "必须关联章节、节次、知识点和证据来源，便于追溯与维护。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4.3 为什么要显式记录禁答点", 2)
    add_para(
        doc,
        "相似问题的错误往往不是完全答非所问，而是混入了邻近概念。仅检查答案是否包含正确内容并不足够，"
        "还应检查是否包含不应出现的内容。例如询问“人机分工的内容”时，答案可以提及具体分工事项，"
        "但不应将“人机分工的原则”作为主要答案。"
    )

    # 5
    add_heading(doc, "5 知识库建库方法", 1)
    add_heading(doc, "5.1 建库总体流程", 2)
    for item in [
        "盘点教材正式稿、知识库初稿、答案卡初稿、图片素材、互动脚本和已有题库。",
        "建立章节结构，明确每章、每节和每个知识点的边界。",
        "从正式稿提炼知识点，再围绕学习问题设计答案卡。",
        "补充同义问法、易混辨析、知识关系、证据片段和练习题。",
        "将图片、公式、表格、互动脚本结构化为可检索资源对象。",
        "建立QA评测集与资源评测集，先评测再发布。",
        "生成嵌入语料、导入数据库、构建向量索引并执行回归测试。",
    ]:
        add_number(doc, item)
    add_heading(doc, "5.2 从章节结构到知识点", 2)
    add_para(
        doc,
        "知识点应围绕可学习、可解释、可关联的最小教学单元设计。过大的知识点难以复用，过小的知识点会造成"
        "答案卡碎片化。通常可从标题层级、定义句、方法步骤、分类体系、对比关系和图表主题中识别知识点。"
    )
    add_heading(doc, "5.3 从知识点到答案卡", 2)
    add_para(
        doc,
        "一个知识点可以产生多个答案卡，因为学生可能从定义、功能、组成、原则、流程、优缺点、适用场景等不同角度提问。"
        "答案卡设计应优先覆盖高频学习意图和易混问题，而不是机械地为每个段落生成问题。"
    )
    add_table(
        doc,
        ["问题类型", "典型问法", "推荐输出模式", "建设要点"],
        [
            ["定义", "什么是……", "paragraph / bullet", "首句给定义，再列关键特征"],
            ["功能", "……有哪些功能", "bullet", "按功能对象或任务分组"],
            ["组成", "……由什么构成", "bullet / table", "明确层级关系"],
            ["原则", "……应遵循哪些原则", "bullet", "原则与具体内容分开"],
            ["流程", "……如何进行", "step", "步骤按先后顺序表达"],
            ["对比", "A与B有什么区别", "table", "统一比较维度，避免散点描述"],
            ["资源学习", "我想看图/公式/互动", "resource_guidance", "直接推荐目标资源并说明用途"],
        ],
        widths=[2.2, 3.8, 3.4, 6.0],
    )
    add_heading(doc, "5.4 同义问法与学生表达", 2)
    add_para(
        doc,
        "教材术语往往规范但不等同于学生问法。同义问法库应覆盖简称、口语表达、疑问句变体、动作词变体、"
        "对象范围词和常见错别字，但不应把语义不同的问题强行归并。高价值问法可赋予更高优先级。"
    )
    add_heading(doc, "5.5 易混问题与对比对象", 2)
    add_para(
        doc,
        "建库时应主动寻找“词面相近、答案不同”的问题对，并为其建立对比关系、范围词、禁答点和测试用例。"
        "本项目中特别关注平面、纵断面、横断面、交叉口、人机分工，以及“功能、内容、原则”等意图差异。"
    )
    add_heading(doc, "5.6 证据片段建设", 2)
    add_para(
        doc,
        "Source_Chunks用于证明答案来自何处。片段应保留章节、节次、原文位置和文本内容，并标记"
        "no_direct_output或等价策略，防止系统把未经整理的原文段落直接当作标准答案。"
    )
    add_heading(doc, "5.7 评测集同步建设", 2)
    add_para(
        doc,
        "答案卡和评测题应同步建设。每张重要答案卡至少应有一个规范问法和若干变体问法；每组易混答案卡应有"
        "对抗性测试；每个资源对象应有能够明确触发它的资源导向问题。"
    )

    # 6
    add_heading(doc, "6 Word正式稿核验与资源抽取", 1)
    add_heading(doc, "6.1 正式稿核验原则", 2)
    add_para(
        doc,
        "知识库初稿可以提高建库效率，但不能替代正式稿核验。核验重点包括术语、定义、要点完整性、章节归属、"
        "图表编号、公式编号、互动脚本内容和来源位置。"
    )
    add_heading(doc, "6.2 Word文本与数学公式读取", 2)
    add_para(
        doc,
        "Word的docx文件本质上是ZIP包。正文文本主要位于document.xml中，复杂公式可能位于Office Math节点。"
        "仅读取普通文本节点会遗漏公式，因此抽取程序应同时遍历Word文本命名空间和数学命名空间，并按文档顺序拼接。"
    )
    add_code_block(
        doc,
        [
            "docx ZIP",
            "  ├─ word/document.xml        正文、段落、表格、数学节点",
            "  ├─ word/_rels/*.rels        图片与关系映射",
            "  └─ word/media/*             图片等媒体文件",
        ],
    )
    add_heading(doc, "6.3 原文片段核验策略", 2)
    add_table(
        doc,
        ["核验状态", "含义", "处理建议"],
        [
            ["exact", "知识库文本可在正式稿中精确找到", "可直接通过"],
            ["partial", "主体内容匹配，但存在格式或局部差异", "人工快速复核"],
            ["concept_supported", "概念得到正式稿支持，但表述不是原文", "检查是否为合理教学归纳"],
            ["unverified", "无法从正式稿确认", "不得直接发布，需人工处理"],
            ["empty", "文本为空或缺失", "补全或删除对象"],
        ],
        widths=[3.2, 6.0, 6.0],
    )
    add_heading(doc, "6.4 图片、表格、公式与互动脚本资源化", 2)
    add_para(
        doc,
        "资源不是答案的装饰物，而是独立的学习对象。每个资源应具有稳定ID、类型、标题、编号别名、路径或结构化内容、"
        "关联知识点、触发问题和状态。"
    )
    add_table(
        doc,
        ["资源类型", "结构化重点", "典型触发问题"],
        [
            ["image", "图号、图题、文件路径、上下文、关联知识点", "我想看图5-1；这张图说明什么"],
            ["formula", "公式号、表达式、变量说明、上下文、关联答案卡", "公式5-1怎么理解"],
            ["table", "表号、标题、列名、行数据、上下文", "表5-1有哪些内容"],
            ["interactive_html", "标题、步骤、入口路径、触发问法、学习目标", "如何操作这个互动；我想体验……"],
        ],
        widths=[3.0, 7.2, 5.0],
    )
    add_heading(doc, "6.5 保守绑定策略", 2)
    add_para(
        doc,
        "自动程序可以依据图号、表号、标题、相邻段落和关联知识点进行绑定，但当编号重复、标题不一致、资源类型冲突或"
        "文档排版复杂时，应降低自动绑定置信度并输出待核验报告。宁可保留少量占位资源，也不要错误地把资源绑定到错误知识点。"
    )

    # 7
    add_heading(doc, "7 数据规范化、校验与版本发布", 1)
    add_heading(doc, "7.1 为什么需要规范化层", 2)
    add_para(
        doc,
        "不同章节、不同编写人员或不同生成工具产生的初稿往往存在表名、字段名、列表格式、ID规则和状态值不一致的问题。"
        "规范化脚本负责把多种输入形态转换为统一的发布模型，使后续检索、评测和数据库导入不依赖章节特例。"
    )
    add_heading(doc, "7.2 规范化的主要任务", 2)
    for item in [
        "将不同表名映射为统一表名。",
        "将字符串、数组、空值等多种字段形态统一为标准列表或结构。",
        "补齐章节ID、节次ID、稳定对象ID和状态字段。",
        "统一答案卡、知识点、资源、同义问法、关系、评测集和嵌入语料结构。",
        "将Source_Chunks标记为证据专用。",
        "将评测题的expected_points和response_mode与目标答案卡对齐。",
        "修正已知重复问题、空ID、占位嵌入文本和错误资源路径。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.3 稳定ID设计", 2)
    add_para(
        doc,
        "稳定ID应在对象生命周期内保持不变，避免数据库更新、向量索引、评测记录和前端链接失效。"
        "当初稿缺少ID时，可以基于章节、对象类型和规范化文本生成哈希ID，但正式发布后不应因文字微调随意更换。"
    )
    add_heading(doc, "7.4 发布前校验规则", 2)
    add_table(
        doc,
        ["校验类别", "示例规则"],
        [
            ["结构完整性", "必需表存在且类型正确"],
            ["对象唯一性", "答案卡、知识点、资源、评测题ID不得重复"],
            ["字段完整性", "答案卡必须有规范问题、答案模式、答案要点、必答点、状态"],
            ["引用一致性", "评测题目标、资源关联知识点、关系目标必须存在"],
            ["内容质量", "答案要点不得为空，正式发布对象状态应为checked或等价状态"],
            ["歧义控制", "重复问题必须有明确区分或显式消歧策略"],
        ],
        widths=[4.0, 11.2],
    )
    add_heading(doc, "7.5 推荐版本发布流程", 2)
    add_code_block(
        doc,
        [
            "raw 初稿与正式稿",
            "  → normalize 规范化",
            "  → validate 结构与引用校验",
            "  → review Word核验与资源报告",
            "  → evaluate QA与资源回归",
            "  → export JSON发布包 / SQL种子",
            "  → embed 生成向量",
            "  → import PostgreSQL / pgvector",
            "  → db-evaluate 数据库模式回归",
            "  → release 版本标记与留档",
        ],
    )

    # 8
    add_heading(doc, "8 检索与RAG编排技术细节", 1)
    add_heading(doc, "8.1 检索顺序", 2)
    add_para(
        doc,
        "本项目采用答案卡优先的混合检索。关键词与规则检索负责稳定识别明确问法和范围词，向量检索负责补充语义相近表达，"
        "知识点、对比对象和证据片段作为后备与解释来源。"
    )
    add_code_block(
        doc,
        [
            "章节上下文识别",
            "  → 问题规范化与分词",
            "  → 资源导向意图判断",
            "  → 答案卡关键词/模式检索",
            "  → 答案卡向量检索",
            "  → 混合评分与消歧重排",
            "  → 知识点/对比对象后备",
            "  → Source_Chunks证据查找",
            "  → answer_mode确定性渲染",
            "  → must_include / avoid 检查",
            "  → 返回答案、资源与检索轨迹",
        ],
    )
    add_heading(doc, "8.2 问题规范化与中文分词", 2)
    add_para(
        doc,
        "问题规范化包括小写化、去空格、去常见标点等处理。为避免依赖重型分词器，本项目同时提取英文/数字词元与中文"
        "2、3、4字gram，用于衡量问题与答案卡检索文本的重叠程度。其他项目可替换为领域分词器，但应保留规范化一致性。"
    )
    add_heading(doc, "8.3 关键词与模式评分", 2)
    add_para(
        doc,
        "检索评分应让精确学生问法和明确问题模式拥有最高权重，其次是规范问题、意图核心、同义问法和范围词匹配；"
        "答案正文只提供较小加分，避免长答案因包含大量通用词而抢占结果。"
    )
    add_table(
        doc,
        ["信号", "作用", "设计原则"],
        [
            ["规范问题精确匹配", "识别标准问法", "高权重"],
            ["学生问题模式精确匹配", "识别真实用户表达", "最高权重之一"],
            ["同义问法匹配", "扩展表达覆盖", "可配置优先级"],
            ["范围词匹配", "区分平面、纵断面、横断面等对象", "显式加权"],
            ["意图词匹配", "区分功能、内容、原则、流程等", "显式加权"],
            ["答案正文重叠", "提供弱语义补充", "低权重，防止长文抢答"],
            ["必答点匹配", "强化核心概念", "中等加权"],
            ["向量相似度", "覆盖未登记的语义变体", "保守增益，不覆盖明确词面信号"],
        ],
        widths=[4.0, 5.0, 6.2],
    )
    add_heading(doc, "8.4 相似问题消歧", 2)
    add_para(
        doc,
        "对相似问题，仅靠向量相似度通常不够。系统应显式识别范围词和意图词，并对冲突候选降权。"
        "例如问题包含“纵断面”和“功能”时，应提升纵断面功能答案卡，降低平面功能、纵断面原则等候选。"
    )
    add_heading(doc, "8.5 连续追问的上下文控制", 2)
    add_para(
        doc,
        "连续追问不应简单沿用上一轮答案。新的问题文本必须拥有主导权，上一轮上下文只用于补全省略信息。"
        "当用户明确输入新的范围词或意图词时，系统应重置相应检索约束，避免上一轮主题污染。"
    )
    add_heading(doc, "8.6 答案渲染", 2)
    add_para(
        doc,
        "答案渲染依据answer_mode完成：bullet用于要点列表，table用于统一维度对比，step用于流程，paragraph用于连贯解释。"
        "这种确定性渲染使答案格式可预测，也使评测能够检查输出模式是否正确。"
    )
    add_heading(doc, "8.7 检索轨迹与可解释性", 2)
    add_para(
        doc,
        "API应返回或记录检索轨迹，包括关键词候选、向量候选、混合评分、选中答案卡、资源推荐、证据片段和选择理由。"
        "检索轨迹是定位误答、调整权重和解释系统行为的关键工程资产。"
    )

    # 9
    add_heading(doc, "9 资源导向问答", 1)
    add_heading(doc, "9.1 资源导向问题的独立性", 2)
    add_para(
        doc,
        "当用户明确表示“我想看图”“公式怎么理解”“打开互动”“表格有哪些内容”时，系统目标不是只返回文字答案，"
        "而是选择正确资源并说明其学习用途。因此资源导向问答应拥有独立检索逻辑、输出模式和评测集。"
    )
    add_heading(doc, "9.2 资源匹配信号", 2)
    for item in [
        "资源标题、资源ID和编号别名。",
        "图号、表号、公式号等引用表达。",
        "触发问题、资源关键词和资源说明文本。",
        "关联知识点、关联答案卡和互动脚本目标。",
        "资源类型意图，例如image、formula、table、interactive_html。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "9.3 同号不同类型冲突", 2)
    add_para(
        doc,
        "教材中可能同时存在图5-1、表5-1、公式5-1。若只匹配数字编号，系统会错误推荐。资源检索必须把类型意图作为"
        "强约束：问题提到“公式”时，图和表候选应被明显降权或排除。"
    )
    add_heading(doc, "9.4 资源状态管理", 2)
    add_table(
        doc,
        ["状态", "含义", "前端处理建议"],
        [
            ["bound", "资源已绑定到文件或可访问入口", "可直接预览或打开"],
            ["structured", "资源已转为结构化内容", "可内联展示并参与检索"],
            ["placeholder", "资源对象存在但尚未完成文件绑定", "可显示说明，不应假装可用"],
            ["unverified", "资源信息尚未核验", "默认不发布或标记待审核"],
        ],
        widths=[3.0, 6.2, 6.0],
    )

    # 10
    add_heading(doc, "10 PostgreSQL、pgvector与嵌入技术", 1)
    add_heading(doc, "10.1 数据库角色", 2)
    add_para(
        doc,
        "PostgreSQL保存运行时主数据、版本、检索字段、资源元数据、评测运行和评测明细。JSONB适合保存数组和原始结构，"
        "高频检索字段则应拆为独立列并建立索引。"
    )
    add_heading(doc, "10.2 推荐索引", 2)
    add_table(
        doc,
        ["索引类型", "适用字段", "用途"],
        [
            ["B-tree", "chapter_id、version_id、source_id、status", "精确过滤与关联查询"],
            ["GIN / JSONB", "related_kp_ids、数组字段", "包含关系查询"],
            ["pg_trgm", "规范问题、标题、别名、检索文本", "模糊匹配与相似文本检索"],
            ["HNSW / vector_cosine_ops", "embedding向量", "高效向量近邻检索"],
        ],
        widths=[4.0, 5.4, 5.8],
    )
    add_heading(doc, "10.3 嵌入语料设计", 2)
    add_para(
        doc,
        "Embedding_Corpus不应直接把所有字段无差别拼接。答案卡向量文本应突出规范问题、学生问法、关键词、必答点和"
        "简洁答案要点；资源向量文本应突出标题、编号、类型、触发问法和学习用途。嵌入文本需要可审阅、可重建、可版本化。"
    )
    add_heading(doc, "10.4 向量检索的定位", 2)
    add_para(
        doc,
        "向量检索用于覆盖未登记的表达变体，而不是替代规则和领域消歧。对教材型精准问答，明确术语、编号、范围词和"
        "问题意图通常比纯语义相似度更可靠。因此向量分数应作为保守增益，并设置阈值。"
    )
    add_heading(doc, "10.5 本地嵌入服务注意事项", 2)
    for item in [
        "固定嵌入模型和向量维度，数据库vector列维度必须一致。",
        "批量生成时控制batch size和请求节奏，避免本地服务过载。",
        "嵌入记录必须有唯一embedding_id、source_type和source_id。",
        "数据库导入后核对语料记录数与非空向量数，及时发现漏嵌入。",
        "版本升级或嵌入文本变化后应重新生成向量并执行数据库模式回归。",
    ]:
        add_bullet(doc, item)

    # 11
    add_heading(doc, "11 自动评测与质量闭环", 1)
    add_heading(doc, "11.1 QA评测指标", 2)
    add_table(
        doc,
        ["指标", "定义", "意义"],
        [
            ["Top1命中率", "第一候选是否为目标答案卡", "衡量直接回答准确性"],
            ["Top3命中率", "前三候选是否包含目标答案卡", "衡量候选召回能力"],
            ["必答点覆盖率", "回答是否覆盖expected_points", "衡量内容完整性"],
            ["禁答点检查", "是否出现should_not_include", "衡量相似概念污染"],
            ["格式匹配", "回答模式是否符合response_mode", "衡量呈现稳定性"],
            ["综合通过率", "以上条件是否同时满足", "衡量可发布质量"],
        ],
        widths=[3.2, 6.2, 6.0],
    )
    add_heading(doc, "11.2 资源评测指标", 2)
    add_para(
        doc,
        "资源评测至少检查：第一推荐资源是否正确、目标资源是否出现在Top3、输出模式是否为resource_guidance、"
        "资源类型是否符合用户意图。"
    )
    add_heading(doc, "11.3 评测集设计方法", 2)
    for item in [
        "覆盖规范问法、口语问法、简称、疑问句变体和省略表达。",
        "为相似答案卡设计对抗性问题，确保范围词和意图词能正确区分。",
        "为每种输出模式提供样例，检查bullet、table、step、paragraph等格式。",
        "为每类资源和同号不同类型资源设计测试。",
        "保留历史失败样例并加入长期回归集，防止问题反复出现。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "11.4 评测驱动迭代", 2)
    add_code_block(
        doc,
        [
            "发现误答",
            "  → 查看检索轨迹和候选分数",
            "  → 判断属于内容缺失、问法缺失、消歧不足、资源冲突或数据错误",
            "  → 优先修正知识对象和检索信号",
            "  → 补充失败样例到评测集",
            "  → 重新执行本地与数据库模式回归",
            "  → 发布新版本",
        ],
    )

    # 12
    add_heading(doc, "12 前端学习体验与预置问题题库", 1)
    add_heading(doc, "12.1 前端不是检索调试器", 2)
    add_para(
        doc,
        "学习前端的首要任务是帮助用户学习，而不是展示所有系统参数。检索模式、轨迹、答案卡详情和评测信息可以提供，"
        "但应放在次级区域，默认界面应突出章节、问题、答案和学习资源。"
    )
    add_heading(doc, "12.2 预置问题题库的价值", 2)
    add_para(
        doc,
        "预置问题可以降低用户提问门槛，帮助学习者发现教材中的重点问题，也能把知识库中已经验证的高质量问题暴露给前端。"
        "题库不是替代自由提问，而是为学习路径提供入口。"
    )
    add_heading(doc, "12.3 题库建设原则", 2)
    for item in [
        "按核心概念、系统功能、设计流程、易混辨析、资源学习等类别组织。",
        "每个答案卡或资源目标默认只保留一个最自然的问题，避免重复。",
        "优先展示少量高价值问题，完整题库按需展开，减少界面拥挤。",
        "问题文本应自然、简洁、可直接点击，不使用内部字段名或生硬模板。",
        "对资源问题优先使用“我想看”“说明了什么”“怎么看”“怎么操作”等表达。",
        "按章节和自然顺序排序，保持学习路径连贯。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "12.4 学习前端应具备的基本能力", 2)
    add_table(
        doc,
        ["能力", "用户价值"],
        [
            ["章节切换与上下文提示", "明确当前学习范围"],
            ["自由提问与预置问题", "兼顾探索与引导"],
            ["答案与置信度", "快速获得结构化解释"],
            ["答案卡详情与来源证据", "理解答案依据并支持复核"],
            ["图片、公式、表格、互动资源", "从文字学习扩展到多种学习方式"],
            ["问题历史与本地反馈", "支持连续学习和问题收集"],
            ["评测运行与失败样例查看", "支持研发和内容维护人员定位问题"],
        ],
        widths=[5.2, 10.0],
    )

    # 13
    add_heading(doc, "13 研发思路与项目推进方法", 1)
    add_heading(doc, "13.1 先做样板章，再复制方法", 2)
    add_para(
        doc,
        "教材知识库建设涉及内容、数据、检索、前端和评测多个环节。直接全书铺开容易把问题放大。推荐先选择一个代表性章节，"
        "打通正式稿核验、答案卡、资源、向量、数据库、API、前端和评测全链路，再将方法复制到其他章节。"
    )
    add_heading(doc, "13.2 从问题驱动研发，而不是功能堆叠", 2)
    add_para(
        doc,
        "每轮研发应围绕真实误答或学习障碍展开。例如，相似系统功能串答促使项目加强范围词消歧；"
        "“内容”与“原则”混淆促使项目显式建模问题意图；前端拥挤和题库重复促使项目建立精选展示与去重规则。"
    )
    add_heading(doc, "13.3 内容工程与软件工程并行", 2)
    add_para(
        doc,
        "检索准确性问题往往同时涉及内容数据和代码策略。研发团队需要把内容编写、数据规范、检索算法、自动评测和前端交互"
        "作为一个整体协作，而不是把误答全部归因于模型或提示词。"
    )
    add_heading(doc, "13.4 每次修复都留下可复用资产", 2)
    for item in [
        "误答修复后增加评测样例。",
        "资源绑定后增加引用别名和触发问题。",
        "格式差异处理后更新规范化脚本。",
        "正式稿核验后保存报告与来源位置。",
        "前端体验优化后更新题库生成规则，而不是只手工修改页面。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "13.5 已验证实践与可复用建议的区分", 2)
    add_para(
        doc,
        "项目文档和评审材料应明确哪些能力已经在当前章节和数据上通过评测，哪些属于面向未来的建议。"
        "这种区分有助于控制预期，也便于其他项目根据自身成熟度选择实施范围。"
    )

    # 14
    add_heading(doc, "14 可复用实施SOP", 1)
    add_heading(doc, "14.1 阶段一：项目准备", 2)
    for item in [
        "确定权威内容来源、章节范围、目标用户和主要学习任务。",
        "盘点现有Word、PDF、JSON、题库、图片、视频、互动脚本等资产。",
        "确定对象ID规则、版本规则、目录结构和发布责任人。",
        "选择样板章节并定义验收指标。",
    ]:
        add_number(doc, item)
    add_heading(doc, "14.2 阶段二：内容建库", 2)
    for item in [
        "建立章节结构与知识点。",
        "设计答案卡和答案模式。",
        "补充学生问法、同义词、易混对比和知识关系。",
        "建立证据片段、练习题和初始评测集。",
        "结构化资源并绑定知识点、答案卡和触发问题。",
    ]:
        add_number(doc, item)
    add_heading(doc, "14.3 阶段三：工程化", 2)
    for item in [
        "编写规范化、校验、导出和数据库迁移脚本。",
        "实现答案卡检索、资源检索、混合检索和确定性渲染。",
        "生成嵌入语料并导入PostgreSQL与pgvector。",
        "实现问答API、检索轨迹和评测记录。",
        "实现学习前端与预置问题题库。",
    ]:
        add_number(doc, item)
    add_heading(doc, "14.4 阶段四：验收与发布", 2)
    for item in [
        "执行结构校验、Word核验、QA评测和资源评测。",
        "检查数据库记录数、向量数、索引和版本状态。",
        "使用真实用户问题进行体验测试并补充回归集。",
        "生成发布包、评测报告、变更说明和遗留问题清单。",
        "上线后持续收集反馈，按版本迭代。",
    ]:
        add_number(doc, item)

    # 15
    add_heading(doc, "15 常见问题、失败模式与经验", 1)
    add_table(
        doc,
        ["问题", "原因", "推荐处理"],
        [
            ["相似系统功能串答", "向量相似度高，范围词权重不足", "显式识别范围词，提升问题模式权重，降低冲突候选"],
            ["“内容”答成“原则”", "意图词未建模或上一轮上下文污染", "对功能、内容、原则、流程建立意图约束和对抗测试"],
            ["原文片段直接成为答案", "把Source_Chunks当作主答案来源", "答案卡优先，片段仅作证据并标记no_direct_output"],
            ["同号图、表、公式互相抢占", "只匹配数字编号", "加入资源类型意图和冲突惩罚"],
            ["Word公式抽取为空", "只读取普通文本节点", "同时读取Office Math节点"],
            ["图片编号或标题错位", "文档排版和关系映射复杂", "保守自动绑定，输出人工核验报告"],
            ["向量记录缺失", "空ID、占位文本、服务过载或导入遗漏", "校验嵌入语料，控制批次，核对记录数与非空向量数"],
            ["数据库存在旧数据", "只upsert未清理已删除对象", "按版本或章节执行一致性清理并回归"],
            ["前端题库拥挤重复", "直接展示全部问题或按文本简单拼接", "每目标一个精选问题，默认少量展示，完整题库按需展开"],
            ["命令行默认版本不一致", "章节、版本参数使用隐式默认值", "发布脚本显式传递chapter_id和version_id"],
        ],
        widths=[4.0, 5.0, 6.2],
    )
    add_note(
        doc,
        "经验判断",
        "当误答出现时，先检查目标知识对象是否清晰、问法是否覆盖、检索信号是否冲突，再考虑调整模型或提示词。",
    )

    # 16
    add_heading(doc, "16 迁移到其他项目的方法", 1)
    add_heading(doc, "16.1 可直接复用的部分", 2)
    for item in [
        "答案卡优先、证据片段支撑、资源对象扩展的分层架构。",
        "知识点、答案卡、同义问法、关系、资源、评测集等核心数据模型。",
        "规范化、校验、嵌入、数据库导入、自动评测和发布流程。",
        "关键词规则与向量检索结合的混合检索思路。",
        "预置问题题库与自由提问并存的学习前端设计。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "16.2 需要按领域调整的部分", 2)
    add_table(
        doc,
        ["调整项", "调整依据"],
        [
            ["问题类型体系", "不同学科可能更重视推导、案例、规范条文或故障诊断"],
            ["答案模式", "数学课程可能需要公式推导，法规项目可能需要条款引用"],
            ["分词与同义词", "领域术语、缩写和符号差异明显"],
            ["资源类型", "可增加视频、仿真、数据集、代码示例、实验步骤等"],
            ["评测标准", "根据答案唯一性、容错范围和业务风险设置"],
            ["权限与审计", "企业或高风险领域可能需要更严格的发布审批和访问控制"],
        ],
        widths=[4.0, 11.2],
    )
    add_heading(doc, "16.3 不同成熟度项目的实施建议", 2)
    add_table(
        doc,
        ["成熟度", "建议范围"],
        [
            ["起步型", "先做一个章节或一个主题，建设答案卡、QA评测和关键词检索"],
            ["成长型", "增加资源层、PostgreSQL、向量检索、前端题库和版本发布"],
            ["规模型", "增加跨章节检索、权限、审计、内容工作流、监控、LMS或业务系统集成"],
        ],
        widths=[3.2, 12.0],
    )

    # 17
    add_heading(doc, "17 当前项目成果与规模", 1)
    add_heading(doc, "17.1 已建设内容规模", 2)
    add_para(
        doc,
        "截至2026年6月3日，项目已完成第1章至第5章的知识对象规范化、答案卡检索、资源检索、"
        "数据库与向量链路、学习前端和自动评测建设。以下统计来自当前data/processed目录。"
    )
    add_table(
        doc,
        ["章节", "答案卡", "知识点", "QA题", "资源", "互动脚本", "资源评测题", "嵌入语料"],
        [
            ["第1章", 112, 74, 443, 14, 10, 65, 186],
            ["第2章", 150, 91, 450, 45, 9, 101, 256],
            ["第3章", 181, 118, 286, 65, 4, 134, 299],
            ["第4章", 254, 135, 404, 21, 8, 52, 389],
            ["第5章", 239, 144, 478, 40, 11, 93, 383],
            ["合计", 936, 562, 2061, 185, 42, 445, 1513],
        ],
        widths=[2.0, 1.7, 1.7, 1.7, 1.5, 1.9, 2.2, 2.0],
    )
    add_heading(doc, "17.2 已验证能力", 2)
    for item in [
        "多章节答案卡优先精准问答。",
        "关键词、本地向量和PostgreSQL/pgvector数据库模式检索。",
        "平面、纵断面、横断面等相似范围问题的显式消歧。",
        "功能、内容、原则等相似问题意图的显式消歧。",
        "图片、公式、表格和互动脚本的资源导向问答。",
        "Word正式稿文本、数学节点、图片和表格的抽取与核验。",
        "QA评测、资源评测、失败样例报告和检索轨迹。",
        "前端预置问题题库、精选展示、去重、资源预览和反馈记录。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "17.3 当前遗留与维护事项", 2)
    add_para(
        doc,
        "当前仍需持续完成部分占位图片资源绑定、复杂公式的视觉排版优化、跨章节检索策略、内容发布工作流和生产级权限审计。"
        "数据库向量记录应在每次发布后核对；当前第1章至第5章嵌入语料与非空向量数量已经全部一致。"
    )

    # 18
    add_heading(doc, "18 后续增强方向", 1)
    for item in [
        "建立面向编写、审核、发布的知识内容工作流，减少脚本与人工操作之间的断点。",
        "扩展跨章节、跨课程检索，同时保留章节上下文优先级和消歧能力。",
        "在规则与向量检索之后引入可解释的轻量重排器，并通过评测验证收益。",
        "增强公式、图表和互动资源的视觉呈现与无障碍支持。",
        "将学习反馈、常见误问和未命中问题转化为建库任务与评测样例。",
        "对接LMS、课程平台或企业知识门户，增加用户、权限、学习记录和审计能力。",
        "建立发布监控，包括命中率、低置信度问题、资源打开率和失败样例趋势。",
    ]:
        add_bullet(doc, item)

    # Appendices
    add_page_break(doc)
    add_heading(doc, "附录A 推荐目录结构", 1)
    add_code_block(
        doc,
        [
            "project/",
            "  ├─ data/",
            "  │   ├─ raw/chXX/              原始知识库初稿",
            "  │   └─ processed/chXX/        规范化知识对象",
            "  ├─ assets/                     图片、互动HTML等资源",
            "  ├─ docs/                       架构、计划、进展与方法文档",
            "  ├─ scripts/                    规范化、核验、资源抽取、发布脚本",
            "  ├─ src/",
            "  │   ├─ kb_rag/                 检索、渲染、API、评测",
            "  │   └─ db/                     数据库迁移与查询",
            "  ├─ output/",
            "  │   └─ chXX_rag_engine/        SQL、向量、报告和中间产物",
            "  └─ web/                        学习前端与问题题库",
        ],
    )

    add_heading(doc, "附录B 核心字段清单", 1)
    add_table(
        doc,
        ["对象", "建议必填字段"],
        [
            ["Answer_Cards", "answer_card_id、chapter_id、section_id、canonical_question、question_type、answer_mode、answer_points、must_include、avoid、status"],
            ["Knowledge_Points", "kp_id、chapter_id、section_id、title、definition、key_points、status"],
            ["QA_Evaluation_Testset", "test_id、question、expected_answer_card_id、expected_points、should_not_include、response_mode"],
            ["Resources", "resource_id、chapter_id、type、title、reference_aliases、related_kp_ids、status"],
            ["Interactive_Scripts", "script_id、title、steps、trigger_questions、related_kp_ids、status"],
            ["Embedding_Corpus", "embedding_id、source_type、source_id、text_for_embedding、model或version"],
        ],
        widths=[4.0, 11.2],
    )

    add_heading(doc, "附录C 发布验收清单", 1)
    checklist = [
        "正式稿版本、知识库版本和数据库版本已明确记录。",
        "所有必需数据表存在，关键ID唯一，引用关系有效。",
        "答案卡有答案要点、必答点、禁答点和输出模式。",
        "重要问题和易混问题已进入QA评测集。",
        "资源对象有类型、标题、编号别名、状态和触发问题。",
        "Word核验报告中的unverified与empty项已处理。",
        "QA评测和资源评测达到项目约定阈值。",
        "数据库记录数、嵌入语料数和非空向量数已核对。",
        "关键词、本地向量和数据库模式均完成回归。",
        "前端在桌面和移动视口完成基本可用性检查。",
        "发布包、SQL、向量、报告和遗留问题清单已留档。",
    ]
    for item in checklist:
        add_bullet(doc, "□ " + item)

    add_heading(doc, "附录D 术语表", 1)
    add_table(
        doc,
        ["术语", "说明"],
        [
            ["RAG", "Retrieval-Augmented Generation，检索增强生成"],
            ["答案卡", "围绕明确学习问题组织的标准答案对象"],
            ["知识点", "可复用、可关联的教学知识单元"],
            ["证据片段", "用于追溯正式文稿依据的原文片段"],
            ["资源导向问答", "以选择和呈现图、表、公式、互动等资源为目标的问答"],
            ["混合检索", "结合关键词、规则、模糊匹配和向量相似度的检索方式"],
            ["pgvector", "PostgreSQL中的向量数据类型与近邻检索扩展"],
            ["回归评测", "在版本变化后重复执行，以发现已有能力是否退化的测试"],
        ],
        widths=[3.6, 11.6],
    )

    add_heading(doc, "结语", 1)
    add_para(
        doc,
        "教材型精准RAG项目的本质，是把专业内容转化为可管理的知识资产，再用检索、评测和交互把这些资产稳定地交付给学习者。"
        "当答案卡、证据、资源、版本和评测形成闭环后，系统就不再只是一次性演示，而具备了持续扩展和长期维护的基础。"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return doc


def validate_docx(doc):
    with ZipFile(OUTPUT, "r") as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError("DOCX ZIP validation failed at: " + bad)
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = required.difference(zf.namelist())
        if missing:
            raise RuntimeError("DOCX missing required files: " + ", ".join(sorted(missing)))

    reopened = Document(OUTPUT)
    if len(reopened.paragraphs) < 100:
        raise RuntimeError("Generated document has too few paragraphs")
    if len(reopened.tables) < 10:
        raise RuntimeError("Generated document has too few tables")

    print(f"output={OUTPUT}")
    print(f"size_bytes={OUTPUT.stat().st_size}")
    print(f"paragraphs={len(reopened.paragraphs)}")
    print(f"tables={len(reopened.tables)}")
    print(f"sections={len(reopened.sections)}")


if __name__ == "__main__":
    generated = build_document()
    validate_docx(generated)
